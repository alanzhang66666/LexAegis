import io
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import pytest
from docx import Document
from fastapi.testclient import TestClient
from lexaegis.api import create_app
from lexaegis.gateway import create_gateway
from lexaegis.worker import process_one, recover
from lexaegis.policy import check_claims
from lexaegis.service import seed_followups
from lexaegis.util import digest, sign
from .conftest import demo_task, approve


def task_state(env, case_id, task_id):
    return next(t for t in env["client"].get(f"/api/v1/cases/{case_id}").json()["tasks"] if t["id"]==task_id)


def test_complete_synthetic_workflow_and_source_confirmation(env):
    case,task=demo_task(env)
    task=approve(env,task)
    client=env["client"]
    assert client.post(f"/api/v1/tasks/{task['id']}/send").status_code==200
    assert process_one(env["app"].state.store,env["root"],env["gate"])
    finished=task_state(env,case["id"],task["id"])
    assert finished["status"]=="complete"
    assert finished["artifact"]["provenance"]=="simulation"
    assert not finished["artifact"]["confirmed"]
    assert client.post(f"/api/v1/tasks/{task['id']}/confirm").status_code==409
    for claim in finished["artifact"]["claims"]:
        assert claim["anchors"]
        assert client.post(f"/api/v1/tasks/{task['id']}/claims/{claim['id']}/review").status_code==200
    assert client.post(f"/api/v1/tasks/{task['id']}/confirm").json()["artifact"]["confirmed"]
    exported=client.get(f"/api/v1/tasks/{task['id']}/export").json()
    assert exported["artifact"]["provenance"]=="simulation"
    assert exported["view"]=="cloud"
    assert finished["artifact"]["unused_fact_ids"] == []
    assert finished["artifact"]["gaps"] == []
    assert exported["artifact"]["unused_fact_ids"] == []
    assert client.get(f"/api/v1/cases/{case['id']}").json()["case"].get("followups") == []
    assert "local_text" not in exported["artifact"]["claims"][0]
    assert client.get(f"/api/v1/tasks/{task['id']}/export?view=local").status_code==409
    assert client.post(f"/api/v1/tasks/{task['id']}/local-export-approval",json={"actor":"复核律师","actor_role":"legal_reviewer"}).status_code==200
    local=client.get(f"/api/v1/tasks/{task['id']}/export?view=local").json()
    assert local["artifact"]["claims"][0]["local_text"]
    assert local["bundle_type"]=="local_review_bundle"


def test_followups_from_unused_create_next_disclosure_task(env):
    case, task = demo_task(env)
    task = approve(env, task)
    client = env["client"]
    assert client.post(f"/api/v1/tasks/{task['id']}/send").status_code == 200
    assert process_one(env["app"].state.store, env["root"], env["gate"])
    finished = task_state(env, case["id"], task["id"])
    unused = [decision["fact_id"] for decision in finished["decisions"] if decision.get("output")][-1:]
    with env["app"].state.store.transaction() as repo:
        current = repo.get(case["id"], "case")
        stored = repo.get(task["id"], "task")
        stored["artifact"]["unused_fact_ids"] = unused
        stored["artifact"]["gaps"] = [{"code": "other", "detail": "模拟缺口", "fact_ids": unused}]
        repo.put("task", stored, case["id"])
        seed_followups(repo, current, stored)
    detail = client.get(f"/api/v1/cases/{case['id']}").json()
    todos = detail["case"]["followups"]
    unused_todo = next(item for item in todos if item["kind"] == "unused_fact")
    gap_todo = next(item for item in todos if item["kind"] == "gap")
    assert unused_todo["status"] == "open" and gap_todo["status"] == "open"
    assert client.post(f"/api/v1/cases/{case['id']}/followups/{gap_todo['id']}/dismiss").status_code == 200
    created = client.post(f"/api/v1/cases/{case['id']}/followups/task", json={"transport": "mock"})
    assert created.status_code == 200
    next_task = created.json()
    assert next_task["followup_fact_ids"] == unused
    assert next_task["status"] == "draft"
    after = client.get(f"/api/v1/cases/{case['id']}").json()["case"]["followups"]
    assert next(item for item in after if item["id"] == unused_todo["id"])["status"] == "addressed"
    assert next(item for item in after if item["id"] == gap_todo["id"])["status"] == "dismissed"


@pytest.mark.parametrize("scenario",["linkage","stop"])
def test_risk_and_stop_examples_block_approval(env,scenario):
    case,task=demo_task(env,scenario)
    result=env["client"].post(f"/api/v1/tasks/{task['id']}/evaluate").json()
    assert result["risk"]["blockers"]
    assert result["risk"]["completion"]=="incomplete"
    assert env["client"].post(f"/api/v1/tasks/{task['id']}/approve",json={"request_hash":result["request_hash"]}).status_code==409
    assert env["client"].post(f"/api/v1/tasks/{task['id']}/stop").json()["status"]=="stopped"


def test_unapproved_and_modified_requests_do_not_send(env):
    case,task=demo_task(env)
    client=env["client"]
    assert client.post(f"/api/v1/tasks/{task['id']}/send").status_code==409
    task=approve(env,task)
    old_hash=task["request_hash"]
    decision=task["decisions"][0]
    assert client.patch(f"/api/v1/tasks/{task['id']}/decisions/{decision['fact_id']}",json={"transform":"delete","necessity":"unnecessary"}).status_code==200
    assert client.post(f"/api/v1/tasks/{task['id']}/send").status_code==409
    assert client.post(f"/api/v1/tasks/{task['id']}/approve",json={"request_hash":old_hash}).status_code==409


def test_concurrent_double_submit_reserves_one_send(env):
    case,task=demo_task(env)
    task=approve(env,task)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda _:env["client"].post(f"/api/v1/tasks/{task['id']}/send").status_code,range(2)))
    assert sorted(results)==[200,409]
    data=env["client"].get(f"/api/v1/cases/{case['id']}").json()
    assert len(data["ledger"])==1
    assert len([j for j in data["jobs"] if j["kind"]=="send"])==1


def test_history_invalidates_second_approval(env):
    case,first=demo_task(env)
    second=env["client"].post(f"/api/v1/cases/{case['id']}/tasks",json={"transport":"mock"}).json()
    first,second=approve(env,first),approve(env,second)
    env["client"].post(f"/api/v1/tasks/{first['id']}/send")
    process_one(env["app"].state.store,env["root"],env["gate"])
    result=env["client"].post(f"/api/v1/tasks/{second['id']}/send")
    assert result.status_code==409
    assert "累计披露" in result.text


def test_timeout_becomes_unknown_and_reconcile_never_posts(env):
    import httpx
    class TimeoutClient:
        def post(self,*args,**kwargs):
            raise httpx.ReadTimeout("test timeout")
        def get(self,*args,**kwargs):
            return httpx.Response(404,request=httpx.Request("GET","http://testserver"))
    case,task=demo_task(env)
    approve(env,task)
    env["client"].post(f"/api/v1/tasks/{task['id']}/send")
    process_one(env["app"].state.store,env["root"],TimeoutClient())
    assert task_state(env,case["id"],task["id"])["status"]=="unknown"
    assert env["client"].post(f"/api/v1/tasks/{task['id']}/send").status_code==409
    env["client"].post(f"/api/v1/tasks/{task['id']}/reconcile")
    process_one(env["app"].state.store,env["root"],TimeoutClient())
    assert task_state(env,case["id"],task["id"])["status"]=="unknown"


def test_restart_recovers_inflight_send_as_unknown(env):
    case,task=demo_task(env)
    approve(env,task)
    env["client"].post(f"/api/v1/tasks/{task['id']}/send")
    with env["app"].state.store.transaction() as repo:
        job=repo.list("job",case["id"])[0]
        job["status"]="running"
        repo.put("job",job,case["id"])
    recover(env["app"].state.store)
    assert task_state(env,case["id"],task["id"])["status"]=="unknown"
    assert not process_one(env["app"].state.store,env["root"],env["gate"])


def test_gateway_hash_auth_and_deduplication(env):
    case,task=demo_task(env)
    task=approve(env,task)
    content={"request":task["request"],"approval":task["approval"]}
    envelope={**content,"signature":sign(content,"unit-test-secret-32-characters-long")}
    result=env["gate"].post("/dispatch",json=envelope)
    assert result.status_code==200
    assert result.json()["status"]=="complete"
    assert env["gate"].post("/dispatch",json=envelope).json()==result.json()
    content["request"]["body"]["messages"][1]["content"]="tampered"
    assert env["gate"].post("/dispatch",json=envelope).status_code==403
    envelope["signature"]=sign(content,"unit-test-secret-32-characters-long")
    assert env["gate"].post("/dispatch",json=envelope).status_code==403
    assert env["gate"].get('/requests/'+task['request']['request_id']).status_code==403


def test_synthetic_mode_cannot_accept_user_material_or_custom_payload(env):
    case,task=demo_task(env)
    client=env["client"]
    assert client.post(f"/api/v1/cases/{case['id']}/documents",files={"file":("private.pdf",b"secret","application/pdf")}).status_code==409
    assert client.post(f"/api/v1/cases/{case['id']}/facts",json={"field":"event","value":"secret","source_ids":["x"],"confirmed":True}).status_code==409
    d=task["decisions"][0]
    assert client.patch(f"/api/v1/tasks/{task['id']}/decisions/{d['fact_id']}",json={"transform":"keep","necessity":"required","output":"private data"}).status_code==409


def formal_material(env):
    client=env["client"]
    case=client.post('/api/v1/cases',json={"name":"正式测试案件"}).json()
    doc=Document();doc.add_paragraph("劳动者林甲主张工资12000元。");table=doc.add_table(rows=1,cols=1);table.cell(0,0).text="工资表证据"
    stream=io.BytesIO();doc.save(stream)
    r=client.post(f"/api/v1/cases/{case['id']}/documents",files={"file":("sample.docx",stream.getvalue(),"application/octet-stream")})
    assert r.status_code==200
    process_one(env["app"].state.store,env["root"],env["gate"])
    data=client.get(f"/api/v1/cases/{case['id']}").json()
    document=data["documents"][0]
    block=document["ir"]["blocks"][0]
    fact=client.post(f"/api/v1/cases/{case['id']}/facts",json={"field":"event","value":block["text"],"source_ids":[block["id"]],"confirmed":True}).json()
    return case,document,fact


def _pending_document(env, case_id, name="pending.txt"):
    uploaded = env["client"].post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": (name, "尚未解析的补充材料。".encode("utf-8"), "text/plain")},
    )
    assert uploaded.status_code == 200, uploaded.text
    return uploaded.json()


def test_empty_ir_does_not_crash_and_blocks_outbound(env):
    case, _document, _fact = formal_material(env)
    client = env["client"]
    _pending_document(env, case["id"])
    detail = client.get(f"/api/v1/cases/{case['id']}")
    assert detail.status_code == 200
    assert any(item.get("ir") is None for item in detail.json()["documents"])
    anchors = client.get(f"/api/v1/cases/{case['id']}/anchors")
    assert anchors.status_code == 200
    task = client.post(f"/api/v1/cases/{case['id']}/tasks", json={"transport": "cloud"})
    assert task.status_code == 200, task.text
    evaluated = client.post(f"/api/v1/tasks/{task.json()['id']}/evaluate")
    assert evaluated.status_code == 200, evaluated.text
    body = evaluated.json()
    assert "材料解析、OCR 或转写复核未完成" in body["risk"]["blockers"]
    assert body["risk"]["approval_allowed"] is False
    approved = client.post(
        f"/api/v1/tasks/{task.json()['id']}/approve",
        json={"request_hash": body["request_hash"]},
    )
    assert approved.status_code == 409


def test_pending_parse_extract_does_not_crash(env):
    case, _document, _fact = formal_material(env)
    client = env["client"]
    settings = client.get("/api/v1/bootstrap").json()["settings"]
    payload = {key: settings[key] for key in (
        "cloud_url", "cloud_model", "edge_url", "edge_model", "deployment_mode", "isolation_verified",
    )}
    payload["edge_model"] = "edge-small"
    assert client.put("/api/v1/settings", json=payload).status_code == 200
    _pending_document(env, case["id"])
    mixed = client.post(f"/api/v1/cases/{case['id']}/extract")
    assert mixed.status_code == 409
    assert "复核" in mixed.text
    pending_only = client.post("/api/v1/cases", json={"name": "仅待解析抽取"}).json()
    _pending_document(env, pending_only["id"], "only-pending.txt")
    result = client.post(f"/api/v1/cases/{pending_only['id']}/extract")
    assert result.status_code == 409
    assert "复核" in result.text


def test_formal_txt_upload_parses_paragraphs(env):
    client = env["client"]
    case = client.post("/api/v1/cases", json={"name": "正式文本案件"}).json()
    uploaded = client.post(
        f"/api/v1/cases/{case['id']}/documents",
        files={"file": ("note.txt", "合成劳动者主张工资12000元。\n\n附件为工资表。".encode("utf-8"), "text/plain")},
    )
    assert uploaded.status_code == 200, uploaded.text
    process_one(env["app"].state.store, env["root"], env["gate"])
    document = client.get(f"/api/v1/cases/{case['id']}").json()["documents"][0]
    assert document["extension"] == ".txt"
    assert len(document["ir"]["blocks"]) == 2
    assert "12000" in document["ir"]["blocks"][0]["text"]


def test_formal_missing_rules_and_attacks_block(env):
    case,doc,fact=formal_material(env)
    client=env["client"]
    assert client.post(f"/api/v1/cases/{case['id']}/tasks",json={"transport":"mock"}).status_code==409
    task=client.post(f"/api/v1/cases/{case['id']}/tasks",json={"transport":"cloud"}).json()
    result=client.post(f"/api/v1/tasks/{task['id']}/evaluate").json()
    assert "适用规则尚未正式发布" in result["risk"]["blockers"]
    assert any("四类" in b for b in result["risk"]["blockers"])
    assert not result["risk"]["approval_allowed"]


def test_local_fact_summary_draft_skips_outbound_gate(env):
    case, _document, _fact = formal_material(env)
    client = env["client"]
    task = client.post(f"/api/v1/cases/{case['id']}/tasks", json={"transport": "local"}).json()
    first = client.post(f"/api/v1/tasks/{task['id']}/evaluate").json()
    blockers = first["risk"]["blockers"]
    assert "适用规则尚未正式发布" not in blockers
    assert not any("隔离" in item for item in blockers)
    assert not any("服务商" in item for item in blockers)
    assert not any("四类" in item for item in blockers)
    assert not any("风险信号权重" in item for item in blockers)
    decision = task["decisions"][0]
    patched = client.patch(
        f"/api/v1/tasks/{task['id']}/decisions/{decision['fact_id']}",
        json={
            "transform": "pseudonymize",
            "necessity": "required",
            "output": "劳动者[[LEX:P0001]]主张工资。",
        },
    )
    assert patched.status_code == 200, patched.text
    ready = client.post(f"/api/v1/tasks/{task['id']}/evaluate")
    assert ready.status_code == 200, ready.text
    assert ready.json()["risk"]["approval_allowed"] is True, ready.text
    assert ready.json()["request"]["target"] == "local://fact-summary-draft"
    approved = client.post(
        f"/api/v1/tasks/{task['id']}/approve",
        json={"request_hash": ready.json()["request_hash"]},
    )
    assert approved.status_code == 200, approved.text
    sent = client.post(f"/api/v1/tasks/{task['id']}/send")
    assert sent.status_code == 200, sent.text
    finished = task_state(env, case["id"], task["id"])
    assert finished["status"] == "complete"
    assert finished["artifact"]["provenance"] == "local_deterministic"
    assert finished["artifact"]["claims"]
    assert not any(job["kind"] == "send" for job in client.get(f"/api/v1/cases/{case['id']}").json()["jobs"])


def test_corrected_doc_retains_old_anchors_and_stales_tasks_after_restart(env):
    case,doc,fact=formal_material(env)
    client=env["client"]
    task=client.post(f"/api/v1/cases/{case['id']}/tasks",json={"transport":"cloud"}).json()
    old_anchor=fact["source_ids"][0]
    corrected=client.put(f"/api/v1/cases/{case['id']}/documents/{doc['id']}/correction",json={"block_id":old_anchor,"text":"校正后工资15000元"})
    assert corrected.status_code==200
    assert corrected.json()["version"]==2
    with TestClient(create_app(env["root"],password="test-restart-password")) as restarted:
        restarted.post('/api/v1/login',json={"password":"test-restart-password"})
        anchors=restarted.get(f"/api/v1/cases/{case['id']}/anchors").json()
        assert any(a["id"]==old_anchor and "12000" in a["text"] for a in anchors)
        data=restarted.get(f"/api/v1/cases/{case['id']}").json()
        assert data["facts"][0]["stale"]
        assert data["tasks"][0]["stale"]
        assert restarted.post(f"/api/v1/tasks/{task['id']}/evaluate").status_code==409


def test_case_mapping_isolated_and_consistent(env):
    case,a=demo_task(env)
    b=env["client"].post(f"/api/v1/cases/{case['id']}/tasks",json={"transport":"mock"}).json()
    other,c=demo_task(env)
    assert a["decisions"][0]["output"]==b["decisions"][0]["output"]
    assert a["decisions"][0]["output"]==c["decisions"][0]["output"]=="[[LEX:P0001]]"
    alias=a["decisions"][0]["output"]
    assert alias in a["decisions"][2]["output"]
    first=env["client"].get(f"/api/v1/cases/{case['id']}").json()["case"]
    second=env["client"].get(f"/api/v1/cases/{other['id']}").json()["case"]
    assert first["mapping_entries"] is not second["mapping_entries"]


def test_hallucinated_numbers_and_missing_sources_never_auto_verify(env):
    case,task=demo_task(env)
    valid=task["decisions"][5]["fact_id"]
    claims=check_claims({"claims":[{"text":"工资999999元","source_ids":[valid]},{"text":"无依据陈述","source_ids":["missing"]}]},task,task["fact_snapshot"])
    assert all(c["status"]=="needs_review" for c in claims)
    assert "日期或金额存在未在引用事实中出现的数字" in claims[0]["issues"]
    assert not claims[1]["anchors"]


def test_settings_accepts_frontend_dual_slots_and_bootstraps_them(env):
    payload={
        "cloud_url":"https://api.openai.com/v1","cloud_model":"primary-model","provider_id":None,
        "cloud_primary":{"url":"https://api.openai.com/v1","model":"primary-model","provider_id":None},
        "cloud_standby":{"url":"https://dashscope.aliyuncs.com/compatible-mode/v1","model":"standby-model","provider_id":None},
        "edge_url":"http://localhost:11434/v1","edge_model":"","deployment_mode":"native","isolation_verified":False,
    }
    payload["disclosure_budget_chars"]=12345
    payload["disclosure_budget_entries"]=3
    saved=env["client"].put("/api/v1/settings",json=payload)
    assert saved.status_code==200,saved.text
    body=saved.json()
    assert body["disclosure_budget_chars"]==12345
    assert body["disclosure_budget_entries"]==3
    assert body["cloud_url"]==body["cloud_primary"]["url"]=="https://api.openai.com/v1"
    assert body["cloud_model"]==body["cloud_primary"]["model"]=="primary-model"
    assert body["cloud_standby"]["model"]=="standby-model"
    echoed=env["client"].get("/api/v1/bootstrap").json()["settings"]
    assert echoed["cloud_primary"]==body["cloud_primary"]
    assert echoed["cloud_standby"]==body["cloud_standby"]
    assert "api_key" not in json.dumps(echoed)
    rejected=env["client"].put("/api/v1/settings",json={**payload,"cloud_primary":{**payload["cloud_primary"],"api_key":"secret"}})
    assert rejected.status_code==422


def test_endpoint_change_invalidates_old_task(env):
    case,task=demo_task(env)
    approve(env,task)
    response=env["client"].put('/api/v1/settings',json={"cloud_url":"https://api.openai.com/v1","cloud_model":"new-model","edge_url":"http://localhost:11434/v1","edge_model":""})
    assert response.status_code==200
    assert env["client"].post(f"/api/v1/tasks/{task['id']}/send").status_code==409


def test_public_edge_endpoint_and_cross_site_calls_rejected(env):
    result=env["client"].put('/api/v1/settings',json={"cloud_url":"https://api.openai.com/v1","cloud_model":"x","edge_url":"https://example.com/v1","edge_model":"x"})
    assert result.status_code==409
    assert env["client"].post('/api/v1/demos',json={"scenario":"success"},headers={"Origin":"https://malicious.example"}).status_code==403
    with TestClient(env["app"]) as unauthenticated:
        assert unauthenticated.get('/api/v1/bootstrap').status_code==401


def test_cloud_mock_provider_and_unknown_provider_error(env):
    class Provider:
        api_key="test-only-key"
        calls=0
        def send(self,target,body):
            self.calls+=1
            f=json.loads(body["messages"][1]["content"])["facts"][0]
            return {"claims":[{"text":f["text"],"source_ids":[f["id"]]}]}
    provider=Provider()
    gateway=TestClient(create_gateway(env["temp"]/"test-provider",secret="unit-test-secret-32-characters-long",provider=provider))
    env["client"].put('/api/v1/settings',json={"cloud_url":"https://api.openai.com/v1","cloud_model":"test-model","edge_url":"http://localhost:11434/v1","edge_model":""})
    case,task=demo_task(env,transport="cloud")
    task=approve(env,task)
    assert task["request"]["dependencies"]["dual_slot"]["primary"]["model"]=="test-model"
    assert not task["request"]["dependencies"]["dual_slot"].get("standby")
    assert task["request"]["target"].endswith("/chat/completions")
    env["client"].post(f"/api/v1/tasks/{task['id']}/send")
    process_one(env["app"].state.store,env["root"],gateway)
    assert provider.calls==1
    assert task_state(env,case["id"],task["id"])["artifact"]["provenance"]=="cloud"
    # This is a test double, not evidence of a live model invocation.
