"""Generate the legal-side request brief as a Word document."""
import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "legal" / "LEX-LG-REQ-2026-09-10-01.docx"
HELPER = Path(__file__).resolve().parent / "make_legal_evidence_docx.py"

spec = importlib.util.spec_from_file_location("cs_brief", HELPER)
cs = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cs)

INK = cs.INK
FOREST = cs.FOREST
MUTED = cs.MUTED


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LEX-LG-REQ-2026-09-10-01  ·  对应 LEX-CS-EVID-2026-09-10-01  ·  法学须给出的内容，非计算机代填  ·  第 ")
    cs.set_run_font(run, name="微软雅黑", size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld1)
    run2._r.append(instr)
    run2._r.append(fld2)
    cs.set_run_font(run2, name="微软雅黑", size=8, color=MUTED)
    run3 = p.add_run(" 页")
    cs.set_run_font(run3, name="微软雅黑", size=8, color=MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    cs.set_cell_margins(section)
    add_footer(section)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("案隐盾法学侧须给出的具体内容")
    cs.set_run_font(r, name="微软雅黑", size=22, bold=True, color=FOREST)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)
    r = sub.add_run("对应计算机侧事实材料  ·  先定口径，再写可落入系统的要求")
    cs.set_run_font(r, name="微软雅黑", size=12, color=MUTED)

    cs.add_table(
        doc,
        ["项", "内容"],
        [
            ["材料编号", "LEX-LG-REQ-2026-09-10-01"],
            ["对应计算机侧材料", "LEX-CS-EVID-2026-09-10-01（docs/legal/LEX-CS-EVID-2026-09-10-01.docx）"],
            ["产品版本", "0.9.0rc0（待发布，非正式发行）"],
            ["日期", "2026-09-10"],
            ["传递路径", "法学侧 → A → 计算机侧 / 对应成员"],
            ["阅读前提", "先读计算机侧第1节流程、第2节服务清单、第5节样例A。未对照这三件，不要先写正式外发或10案验收条文。"],
            ["性质", "本文列出法学必须给出的判断、填写栏和签署项。计算机不代作法律结论。吃不准就填「未知」或「暂不开放」，不要留空充作已审。"],
        ],
        col_cm=[3.4, 12.8],
    )
    cs.add_note(
        doc,
        "浅底单元格表示当前仓库仍为未知、待审或未做。"
        "labor_v2 仍是 draft，dispatch_allowed=false。演示规则非法源。"
        "工程基线红队 resisted 不能当作法学攻击结论。",
    )

    cs.add_heading(doc, "〇、建议交付顺序", 1)
    cs.add_p(
        doc,
        "与计算机侧建议对齐：先根据已有对象出口径，再写能进系统的规则，最后才谈正式发布和10案。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["批次", "法学先交什么", "必须已看到的计算机材料", "交完之后计算机才能做什么"],
        [
            [
                "第一批（立即）",
                "合成材料使用口径；本地/模拟/边侧/云端四条路径的用途判断；主备服务商处理条件表（查不到就标未知）；演示01/02/03验收表各填一份",
                "流程表、服务清单、样例A外发全文、阻断样例、校核能力表",
                "把口径写进产品说明；补服务商档案字段；按验收表修提示或门禁文案，不擅自改法源",
            ],
            [
                "第二批（口径稳定后）",
                "冻结一个子场景（建议：事实摘要中的加班核对）；条件规则与精度；处理依据按主体类别；保密与禁止外发清单；法源快照与条号映射",
                "V2草稿包 packs/labor_v2_draft.json；计算机侧第6至8节",
                "把规则从 uncertain 改成可执行值；加班条件规则从待审草稿升为 reviewed 候选，仍不自动开通外发",
            ],
            [
                "第三批（正式外发前）",
                "规则包发布签署；主备档案复核与影响评估；正式攻击者模型与阈值；真实材料授权；10案校准方案",
                "真云一次取证（目前没有）；隔离与账号合同快照",
                "才可能把 pack 标 published、dispatch_allowed=true。现在不要预填这一栏为「允许」",
            ],
        ],
        col_cm=[2.6, 5.0, 4.4, 4.2],
    )

    cs.add_heading(doc, "一、支撑「案例使用权限与用途判断」", 1)
    cs.add_p(
        doc,
        "计算机已提供事实。法学须对具体对象作允许/附条件/禁止，并写清依据与有效期。不能只写原则句。",
        first_line=False,
    )

    cs.add_heading(doc, "1. 合成材料使用口径", 2)
    cs.add_p(
        doc,
        "对象已固定：演示01/02/03、examples/ 合成附件、data/mock_labor_v1 八条、ui-walk 截图。"
        "这些材料声明不含真实案件。法学须判断的是「在本项目中可以怎么用」，不是「它们是不是真案」。",
    )
    cs.add_table(
        doc,
        ["对象", "计算机已写明的用途", "法学须填写：内部研发", "教学演示", "截图/视频", "对外交付/论文", "依据或未知"],
        [
            ["演示01 工资争议（模拟）", "工作台跑通；mock 不打云端", "", "", "", "", ""],
            ["演示02 组合线索", "展示批准409", "", "", "", "", ""],
            ["演示03 主动停止", "展示材料不完整阻断", "", "", "", "", ""],
            ["examples 合成 DOCX/扫描PDF", "正式案解析与409体验", "", "", "", "", ""],
            ["mock_labor 8条", "评测冒烟；dispatch_allowed=false；标注 pending", "", "", "", "", ""],
            ["ui-walk 截图（含合成姓名、token、本机验收口令界面）", "本机目录，不进git，未进CI", "", "", "", "", ""],
            ["边侧 probe-edge 报告", "工程五项，无法学分，写在被忽略目录", "", "", "", "", ""],
        ],
        col_cm=[2.6, 3.2, 2.1, 2.1, 2.1, 2.1, 2.0],
    )
    cs.add_note(doc, "填写用：允许 / 附条件 / 禁止。附条件须写清条件（例如打码姓名、不得含口令、仅队内）。")

    cs.add_heading(doc, "2. 四条处理路径的用途判断", 2)
    cs.add_p(
        doc,
        "同一合成案会走不同路径，外发内容不同。请逐条判断，不要用「AI辅助办案」一笔带过。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["路径", "实际离开边侧的内容（计算机已证）", "拟认定用途", "是否允许（填）", "处理依据（填）", "必要性说明（填）", "告知/同意（填）"],
        [
            [
                "本地解析/OCR/人工事实",
                "不离开本机。原文在 .data/app/files 与 sqlite",
                "办案材料整理（本机）",
                "",
                "",
                "",
                "",
            ],
            [
                "边侧抽取 MiniCPM",
                "原文块全文发给本机 127.0.0.1:8080。演示不走此路",
                "本机抽取候选事实",
                "",
                "",
                "",
                "",
            ],
            [
                "本地摘要 / 模拟摘要",
                "不打公网。披露句只在本机网关拼回",
                "演示或本机草稿",
                "",
                "",
                "",
                "",
            ],
            [
                "云端摘要 DeepSeek 主 / 百炼备",
                "仅 wire：token化披露句 + 提示词。正式门禁关闭，无成功记录",
                "可选云端生成摘要",
                "",
                "",
                "",
                "",
            ],
        ],
        col_cm=[2.4, 3.6, 2.4, 1.8, 2.0, 2.0, 2.0],
    )
    cs.add_p(
        doc,
        "须分别写清：劳动者、用人单位联系人、证人、家属是否适用同一依据。"
        "「履行合同所必需」不能自动覆盖对方和证人，也不能自动证明可选云端调用必要。"
        "律师点击批准不替代个人有效同意。敏感信息须单独判断，不能一律必须同意，也不能一律以办案必要免除。",
    )
    cs.add_p(
        doc,
        "去标识化与匿名化不同：系统保留边侧映射，不能仅因替换成 [[LEX:P0001]] 就排除个人信息规则。"
        "律师法第41条（2026-09-01文本；旧第38条）覆盖国家秘密、商业秘密、当事人隐私及其他不愿泄露的信息，不限于直接识别自然人。",
    )

    cs.add_heading(doc, "3. 服务商处理条件表（主备分填）", 2)
    cs.add_p(
        doc,
        "计算机只收集了公开网页，没有本团队合同、账号类型和控制台开关。"
        "系统复核服务商档案时，下列字段缺一不可：角色、地区、留存、删除机制、训练使用、人工访问、分包商、合同引用、影响评估、审核人、有效期、允许任务/案由/字段/传输。"
        "请按公开条款 + 本团队资料填写；查不到写「未知」，未知则正式外发保持关闭。",
        first_line=False,
    )
    for slot, name, endpoint in (
        ("主槽", "DeepSeek 开放平台 / deepseek-v4-flash", "https://api.deepseek.com/v1/chat/completions"),
        ("备槽", "阿里云百炼 / qwen3.8-flash", "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"),
    ):
        cs.add_heading(doc, f"{slot}：{name}", 3)
        cs.add_table(
            doc,
            ["档案字段", "系统枚举或要求", "计算机已见公开信息", "法学填写"],
            [
                ["法律主体 / 产品", "自由文本", endpoint, ""],
                ["处理角色 role", "processor / independent_controller / other（不可 unknown）", "协议把开发者写成下游处理者；服务商角色不能只凭「API」认定", ""],
                ["处理地区 processing_regions", "须非空。跨境另判", "DeepSeek 隐私政策写在中国存储；百炼默认中国站，国际站可选", ""],
                ["训练使用 training_use", "forbidden / allowed（不可 unknown）", "开放平台协议未写明训练授权；消费者隐私政策写训练且声明不覆盖下游终端用户。百炼帮助页称不用于训练", ""],
                ["人工访问 human_access", "forbidden / controlled / allowed（不可 unknown）", "均提到执法/安全/服务商；是否读 API 正文未知", ""],
                ["留存 retention", "须非空，写期限或「按合同第×条」", "DeepSeek：账号存续+网安法日志六月。百炼：调用数据按法规存储，正文期限未知", ""],
                ["删除机制 deletion_mechanism", "须非空", "邮箱申请或控制台；本系统不能代删对方日志", ""],
                ["分包商 subprocessor_status", "none 或 declared（declared 须列名单）", "类别提及，无具名清单", ""],
                ["合同 contract_reference + version", "须有引用", "公开协议链接已列；本账号点击版本未知", ""],
                ["影响评估 impact_assessment_reference", "须有引用", "仓库无", ""],
                ["有效期 valid_until", "建议填写", "无", ""],
                ["允许任务/案由/字段", "须覆盖拟外发集合", "拟用 fact_summary + labor + 披露字段", ""],
                ["境外路线", "建议：未经核验则不可选", "白名单含 dashscope-intl 与 api.openai.com，不是默认双槽", ""],
            ],
            col_cm=[3.6, 3.8, 5.0, 3.8],
        )

    cs.add_heading(doc, "4. 保存、访问与删除的法学要求", 2)
    cs.add_p(
        doc,
        "计算机已实现：无整案删除、无 TTL、已登录即可下载原件、备份含原件和口令哈希、业务库保存冻结请求全文。"
        "影响评估与处理记录依法至少保存三年，不等于原文、临时载荷和供应商日志全部保存三年。请分开写。",
    )
    cs.add_table(
        doc,
        ["对象", "现状（计算机）", "法学要求保存多久（填）", "谁可访问（填）", "删除条件（填）", "是否须先改系统（填）"],
        [
            ["原件 files/", "无TTL，无删除API，登录即可下", "", "", "", ""],
            ["映射与 local_text", "只在边侧库", "", "", "", ""],
            ["冻结外发稿 body", "在 app.sqlite 披露账本", "", "", "", ""],
            ["审计事件", "无TTL，可导出", "", "", "", ""],
            ["备份 zip", "操作员自定", "", "", "", ""],
            ["服务商侧副本", "系统不能代删", "", "", "", ""],
            ["影响评估/处理记录", "仓库尚未建档", "", "", "", ""],
        ],
        col_cm=[2.6, 3.6, 2.6, 2.4, 2.5, 2.5],
    )

    cs.add_heading(doc, "二、支撑「法律输出验收标准」", 1)
    cs.add_p(
        doc,
        "请对着系统实际会输出的东西写标准。不要要求系统去检查它目前既不计算、也不展示的项目，除非同时标明「列为下一期工程」。",
        first_line=False,
    )

    cs.add_heading(doc, "5. 先冻结一个子场景", 2)
    cs.add_p(
        doc,
        "仓库 V2 草稿只把「核对加班」写成待审条件规则；一般事实摘要字段必要性全是 uncertain；时间线/证据/争点待确认。"
        "争点提取只是按字段分桶，非法学结论。请先选一个子场景签字冻结，不要同时审四类任务。",
    )
    cs.add_table(
        doc,
        ["候选项", "系统现状", "法学选择（择一）", "选择理由（填）"],
        [
            ["劳动争议事实摘要 · 一般整理", "演示包：假名化+按月日期+金额保持。V2：uncertain，未匹配条件则删除", "", ""],
            ["劳动争议事实摘要 · 核对加班（建议优先）", "有条件规则草稿：须留日历日、日性质、小时、基数、主张/否认、证据持有人；禁止默认截月或平移日历", "", ""],
            ["解除经过整理", "无独立规则，仅研究建议", "", ""],
            ["时间线 / 证据清单 / 争点", "本地筛选或分桶，无法律推理", "", ""],
        ],
        col_cm=[4.0, 6.2, 2.6, 3.4],
    )

    cs.add_heading(doc, "6. 条件规则与精度（对着 V2 草稿改，不要另起无法入库的格式）", 2)
    cs.add_p(
        doc,
        "每条规则最少写清：任务及子目的、适用条件、受影响字段、必要性、最低精度、允许变换、禁止默认变换、不可破坏关系、法源编号、缺项策略、条件未知策略、审核人。"
        "系统已能表达这些字段。条件未知必须进复核，不能默许外发。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["规则要点", "草稿现状", "法学裁定（填）", "法源条号与文本日期（填）"],
        [
            ["一般摘要要不要披露主体/关系/事件", "uncertain；演示包则假名化后可发", "", ""],
            ["一般摘要日期精度", "演示包按月泛化。加班草稿禁止截月", "", ""],
            ["一般摘要金额", "演示包 keep", "", ""],
            ["加班：日历日 / 日性质 / 小时 / 基数", "conditional，preferred=原精度，待审", "", ""],
            ["加班：主张或否认不得写成无争议事实", "草稿已写，待审", "", ""],
            ["加班：谁持有证据", "草稿 tokenize，待审", "", ""],
            ["病历/工伤全文", "代码无专用规则；研究建议全文留本地", "", ""],
            ["混同用工多公司", "代码无专用规则；研究建议分别令牌化", "", ""],
            ["禁止外发词/组合线索", "演示禁止「唯一」「公开获奖」。正式清单空", "", ""],
        ],
        col_cm=[4.2, 4.6, 3.6, 3.8],
    )
    cs.add_note(
        doc,
        "法源须附 2026-09-01 律师法第41条快照及旧第38条映射；加班举证责任勿把解释（二）当现行入口，解释（二）只处理适用时间。"
        "草稿摘述不是正式公布文本，发布前须换成核对过的条文快照。",
    )

    cs.add_heading(doc, "7. 输出对错标准（须能打在验收表上）", 2)
    cs.add_p(
        doc,
        "计算机已会：引用是否存在、数字是否在引用句中出现、token 是否破损、直接标识正则、改写只提示不整单作废。"
        "计算机不会：主张写成认定、主体对调、证据挂错人、日期过了除斥/时效边界的法律含义。请把「必须改」和「建议改」分开，并写清是否要求本期就工程化。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["评价项", "系统现在", "错误等级（填：阻断/必须改/建议改/可接受）", "是否本期写入规则（填）", "说明（填）"],
        [
            ["引用不存在或指向未披露事实", "已阻断确认", "", "", ""],
            ["摘要数字未在引用事实中出现", "已阻断确认", "", "", ""],
            ["改写或新增陈述", "只标 needs_review", "", "", ""],
            ["把「主张」写成「加班了」", "未检查", "", "", ""],
            ["劳动者与用人单位对调", "未检查（硬约束草稿有，未当摘要校核）", "", "", ""],
            ["否定/抗辩丢失", "未检查", "", "", ""],
            ["证据持有人改挂", "未检查", "", "", ""],
            ["日期截月导致期限错误", "演示包会截月；加班草稿禁止", "", "", ""],
            ["间接识别（唯一岗位+公开事件）", "演示禁止词；正式无清单", "", "", ""],
            ["读起来通顺但缺要件", "不检查「通顺」", "", "", ""],
        ],
        col_cm=[3.6, 3.8, 3.2, 2.8, 2.8],
    )
    cs.add_p(
        doc,
        "效用标准请写成「能否回答本子场景的任务问题」，不要写成「摘要是否流畅」。"
        "被阻断和放弃的任务也要进验收分母，避免只评成功跑通的案子。",
    )

    cs.add_heading(doc, "8. 演示三案验收表（第一批就要填）", 2)
    cs.add_p(
        doc,
        "输出来源必须标明：演示01是模拟拼接，不是模型生成。02/03是阻断，没有摘要。不要用这三案给云端模型打质量分。",
        first_line=False,
    )
    for case_id, source, note in (
        ("DEMO-01", "模拟拼接 · labor 0.1.0-demo", "七条披露：两枚 token、按月日期、金额12000；local_text 还原林澄/星桥"),
        ("DEMO-02", "无摘要 · 批准409", "组合线索：唯一岗位+公开获奖"),
        ("DEMO-03", "无摘要 · 主动停止", "关键证据不完整"),
    ):
        cs.add_heading(doc, f"{case_id}  {note}", 3)
        cs.add_table(
            doc,
            ["评价项", "输出来源", "错误位置", "错误等级", "修改建议", "复核人", "结论"],
            [
                ["主体关系是否仍成立", source, "", "", "", "", ""],
                ["日期精度是否可接受", source, "", "", "", "", ""],
                ["金额与主张表述", source, "", "", "", "", ""],
                ["证据归属", source, "", "", "", "", ""],
                ["不应披露却仍留下的线索", source, "", "", "", "", ""],
                ["阻断原因是否成立（02/03）", source, "", "", "", "", ""],
            ],
            col_cm=[3.2, 3.0, 2.0, 2.0, 2.4, 1.8, 1.8],
        )

    cs.add_heading(doc, "9. 标注与10案校准", 2)
    cs.add_p(
        doc,
        "8条 mock_labor 双人标注仍为 pending。没有标注不得生成法律可用性分数。"
        "10案校准须等：第一批口径、一个冻结子场景、以及计算机补齐对应输出（目前没有10份真实模型摘要）。"
        "半合成晋级后仍须 dispatch_allowed=false。真实文书不得入库。",
    )
    cs.add_table(
        doc,
        ["项", "最少写清", "法学填写"],
        [
            ["标注人 A/B 与仲裁人", "具名，三人不同", ""],
            ["先标哪几条", "建议先 DEMO-01 对照 + mock_labor 开发集，不要先盲测", ""],
            ["必要性取值", "required / conditional / unnecessary / uncertain", ""],
            ["变换取值", "keep / pseudonymize / generalize / delete", ""],
            ["10案何时开始", "须写触发条件，不要先定日期", ""],
            ["真实案卷是否进入试点", "须组织授权；授权范围、主体、期限", ""],
        ],
        col_cm=[3.6, 7.0, 5.6],
    )

    cs.add_heading(doc, "三、正式发布与门禁（第三批，现在不要勾「允许外发」）", 1)
    cs.add_p(
        doc,
        "系统发布案由包要求：status 先到 reviewed；有法源 references、审核人、审核记录；legal_status=reviewed；"
        "dispatch_allowed；published_scope；五类风险权重 entity/density/context/combination/cross_round 全部冻结。"
        "另外还要：主备服务商均 reviewed 且未过期；四类攻击达到法学承认的 resisted（不是空跑工程基线）；隔离自检通过。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["签署项", "系统卡点", "法学填写：结论", "版本/日期", "签署人"],
        [
            ["V2 规则包是否可标 reviewed", "现为 draft.1，加班规则待审", "", "", ""],
            ["是否允许 dispatch_allowed=true", "现为 false", "", "", ""],
            ["published_scope（适用法院/时间/任务）", "空", "", "", ""],
            ["五类风险权重", "未冻结则正式云端阻断", "", "", ""],
            ["主槽档案复核", "缺合同与评估即阻断", "", "", ""],
            ["备槽档案复核", "未审核则不得纳入批准目标", "", "", ""],
            ["正式攻击者模型与阈值", "工程 resisted ≠ 法学背书", "", "", ""],
            ["真实材料组织授权", "仓库禁止真实案入库", "", "", ""],
        ],
        col_cm=[3.8, 4.2, 3.0, 2.6, 2.6],
    )

    cs.add_heading(doc, "四、不要写的内容", 1)
    cs.add_table(
        doc,
        ["容易提前写出来、但系统对不上的要求", "为什么现在不要写死"],
        [
            ["「云端摘要必须达到某准确率」", "没有真实云端成功记录，也没有10份模型输出"],
            ["「四类攻击已通过即可外发」", "当前执行器是工程基线；空辅助不得当成 resisted 的法律含义"],
            ["「争点提取须识别法律关系」", "现实现只是六类字段分桶"],
            ["「日期一律按月即可」", "与加班草稿和期限边界冲突，须按子目的分别写"],
            ["「已假名化即不适用个人信息规则」", "边侧保留映射，不是匿名化"],
            ["「律师批准即处理合法」", "批准是发送控制，不是第13条依据本身"],
            ["把2017年律师法第38条当作现行条号", "2026-09-01 起对应第41条"],
        ],
        col_cm=[6.4, 9.8],
    )

    cs.add_heading(doc, "五、回传清单（交 A 时勾选）", 1)
    cs.add_table(
        doc,
        ["勾选", "材料", "批次", "对应计算机材料"],
        [
            ["□", "合成材料使用口径表（第1节）", "第一批", "CS 第5节样例与公开包"],
            ["□", "四条路径用途判断表（第2节）", "第一批", "CS 第1节流程"],
            ["□", "主槽处理条件表（第3节）", "第一批", "CS 第2至3节 DeepSeek"],
            ["□", "备槽处理条件表（第3节）", "第一批", "CS 第2至3节 百炼"],
            ["□", "保存访问删除要求（第4节）", "第一批", "CS 第4节"],
            ["□", "DEMO-01/02/03 验收表（第8节）", "第一批", "CS 第8至9节"],
            ["□", "子场景选择（第5节）", "第二批", "labor_v2_draft.json"],
            ["□", "条件规则与精度裁定（第6节）", "第二批", "V2 矩阵与 disclosure_rules"],
            ["□", "输出对错标准（第7节）", "第二批", "CS 校核能力表"],
            ["□", "标注人与10案触发条件（第9节）", "第二批", "data/ANNOTATION.md"],
            ["□", "发布与门禁签署（第三节）", "第三批", "governance.publish_pack_record"],
        ],
        col_cm=[1.6, 5.4, 2.4, 6.8],
    )

    end = doc.add_paragraph()
    end.paragraph_format.space_before = Pt(18)
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run("（全文完。本文是法学须交付内容的清单与填空，不构成已完成审核或允许外发。）")
    cs.set_run_font(r, name="微软雅黑", size=10, color=MUTED)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
