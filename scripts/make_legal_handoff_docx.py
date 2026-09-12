"""Generate the 2026-09-12 legal handoff brief from the 10-case pilot. No case text."""
import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "legal" / "LEX-LG-HANDOFF-2026-09-12-01.docx"
HELPER = Path(__file__).resolve().parent / "make_legal_evidence_docx.py"

spec = importlib.util.spec_from_file_location("cs_brief", HELPER)
cs = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cs)

INK = cs.INK
FOREST = cs.FOREST
MUTED = cs.MUTED
DOC_ID = "LEX-LG-HANDOFF-2026-09-12-01"


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{DOC_ID}  ·  0.9.0rc0  ·  试点后法学交接，非法学结论、非外发开通  ·  第 ")
    cs.set_run_font(run, name="微软雅黑", size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld1)
    run2._r.append(instr)
    run2._r.append(fld2)
    cs.set_run_font(run2, name="微软雅黑", size=8, color=MUTED)
    run3 = p.add_run(" 页")
    cs.set_run_font(run3, name="微软雅黑", size=8, color=MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    cs.set_cell_margins(section)
    add_footer(section)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = cs.Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = cs.Pt(8)
    title.paragraph_format.space_after = cs.Pt(4)
    r = title.add_run("案隐盾试点后须交付法学的资料与待裁定项")
    cs.set_run_font(r, name="微软雅黑", size=20, bold=True, color=FOREST)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = cs.Pt(16)
    r = sub.add_run("10 案层 A 评估已入账  ·  0 案真实发送  ·  计算机不代填法律结论")
    cs.set_run_font(r, name="微软雅黑", size=12, color=MUTED)

    cs.add_table(
        doc,
        ["项", "内容"],
        [
            ["材料编号", DOC_ID],
            ["日期", "2026-09-12"],
            ["产品版本", "0.9.0rc0（待发布，非正式发行）"],
            ["对应先发材料", "LEX-CS-EVID-2026-09-10-01；LEX-LG-REQ-2026-09-10-01（其中空白栏仍待填）"],
            ["试跑目录", "本机 gitignore：.data/real_cases/runs/pilot-2026-09/（真案正文不入库、不随本文附送）"],
            ["传递路径", "计算机侧 → A → 法学侧。签署回传后再谈层 B 发送"],
            ["性质", "本文说明：工程已交出什么公开计数、法学现在必须裁定什么。浅底格为待填。吃不准填「未知」或「暂不开放」，不要留空充作已审。"],
            ["禁止写法", "不得写成「10 案已通过」或「产品外发已开通」。可写：10 案评估完成，0 案经授权真实发送。"],
        ],
        col_cm=[3.4, 12.8],
    )
    cs.add_note(
        doc,
        "本文不含当事人姓名、裁判原文、清单具体内容或云端密钥。"
        "案号 LA-00x / SMP-00x 是试点稳定编号，不是法院案号。",
    )

    cs.add_heading(doc, "一、当前资料情况（工程已交出）", 1)
    cs.add_p(
        doc,
        "试点任务固定为「事实摘要辅助与披露控制」。正式外发门禁未改。仓库 V2 草稿仍为 draft。native 不能自证隔离。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["资料", "状态", "法学现在怎么用"],
        [
            ["10 案分层与编号（LA-001–005，SMP-001–005）", "已整理。文书层与参考摘要已拆开", "只评文书层；参考摘要不得当原文"],
            ["资料/规则覆盖表与待裁定清单", "已编入试点 coverage/", "补裁定，不要改工程凑绿"],
            ["标注指南草稿 G1–G5", "draft-from-cal.md，未冻结、未签署", "逐条裁定后才能定版本号"],
            ["CAL-1/2/3（LA-001、LA-003、SMP-002）", "Path A/B + 层 A 完成", "先交叉核对这三案 Path B"],
            ["其余 7 案", "Path A/B + 层 A 完成（指南未冻）", "不宣称盲测"],
            ["层 A 评估包", "10 案均有 request_hash、blocker 全文、本地确定性摘要", "409 是有效结果"],
            ["逐案授权表", "仍为空", "必须逐行填写，见第四节"],
            ["试点案由包 / 主备档案 / 攻击阈值", "未签署、未发布", "见第五节；未签则停在层 A"],
            ["Path C 十条边界", "未跑", "须先确认预期，见第六节"],
            ["层 B 发送账本", "无 Attempt。CAL-1 曾探 send，亦 409", "未授权、门禁未过，不得发送"],
        ],
        col_cm=[5.2, 5.6, 5.4],
    )

    cs.add_heading(doc, "二、10 案公开计数（不含原文）", 1)
    cs.add_p(
        doc,
        "Path A 机械确认规则：抽取 value 是文书层连续子串，且不少于 2 字。未确认抽取不得标成模型自动完成。"
        "Path B 只写入清单「具体内容」能对上文书层、且类型能映射到案由字段的行。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["编号", "角色", "抽取", "机械确认", "Path B 写入（未映射 / 不在文书层）", "A2 / B1", "批准"],
        [
            ["LA-001", "CAL-1", "77", "72（跳过 4，短 1）", "7（校准轮）", "blocked / blocked", "409"],
            ["LA-003", "CAL-2", "91", "65（跳过 21，短 5）", "7（校准轮）", "blocked / blocked", "409"],
            ["SMP-002", "CAL-3 工伤含表", "34", "28（跳过 6）", "11（校准轮）", "blocked / blocked", "409"],
            ["LA-002", "后 7", "38", "27（跳过 7，短 4）", "7（13 / 16）", "blocked / blocked", "409"],
            ["LA-004", "后 7", "64", "53（跳过 9，短 2）", "1（8 / 6）", "blocked / blocked", "409"],
            ["LA-005", "后 7", "68", "51（跳过 13，短 4）", "2（6 / 11）", "blocked / blocked", "409"],
            ["SMP-001", "后 7", "56", "44（跳过 12）", "12（20 / 24）", "blocked / blocked", "409"],
            ["SMP-003", "后 7", "48", "31（跳过 17）", "5（24 / 16）", "blocked / blocked", "409"],
            ["SMP-004", "后 7", "82", "50（跳过 31，短 1）", "10（23 / 18）", "blocked / blocked", "409"],
            ["SMP-005", "后 7", "62", "54（跳过 6，短 2）", "8（17 / 9）", "blocked / blocked", "409"],
        ],
        col_cm=[2.0, 2.2, 1.4, 3.0, 3.4, 2.4, 1.8],
    )
    cs.add_note(
        doc,
        "共同 blocker：适用规则未正式发布；V2 包未冻结；服务商处理条件或影响评估不完整；"
        "四类冻结攻击未全部 resisted；隔离自检未通过；五类风险权重未冻结。"
        "Path A 另有：仍包含未替换的登记主体；任务中仍含未确认事实。",
    )

    cs.add_heading(doc, "三、标注指南草稿：须裁定后才能冻结", 1)
    cs.add_p(
        doc,
        "下列 G1–G5 来自校准证据，已回链 issue 类型。未冻结不得层 B 发送。请每条选一个可执行选项并署名。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["条", "工程证据（无原文）", "请裁定（择一）", "依据 / 有效期", "签署"],
        [
            [
                "G1 文书层与清单必须同指一块原文",
                "Path B：清单「具体内容」不在文书层，校准三案约 17+6+18 行；后 7 案同样大量「不在层」",
                "□ 标错层，清单应改\n□ 允许「仅确认事实、无判决原文」外发\n□ 暂不开放\n□ 未知",
                "",
                "",
            ],
            [
                "G2 清单类型要落到案由字段",
                "校准三案约 18+11+29 行无法映射到 entity/relation/event/date/amount/evidence",
                "□ 身份证/手机等禁止进事实表\n□ 只作主体属性假名化\n□ 禁止外发\n□ 未知",
                "",
                "",
            ],
            [
                "G3 边侧抽取是未确认草稿",
                "截断 JSON 已修好。抽取可入库，但 event 仍偏多、偶发 1 字 value。分母只用机械确认事实",
                "□ 同意：未确认不得当已校对\n□ 须人工双人确认才进分母\n□ 未知",
                "",
                "",
            ],
            [
                "G4 缺省披露会清空 wire",
                "B0 缺省 uncertain=delete，三案 0 条 wire。试验策略才有 keep/假名化",
                "□ 确认缺省删除只是保守工程，非法源方案\n□ 另行给出缺省矩阵\n□ 未知",
                "",
                "",
            ],
            [
                "G5 保留原文主体会触发登记名阻断",
                "keep 非主体字段后仍可能报「仍包含未替换的登记主体」",
                "□ 主体一律假名化\n□ 允许 keep 但须先去登记名\n□ 未知",
                "",
                "",
            ],
        ],
        col_cm=[3.2, 4.6, 4.4, 2.2, 1.8],
    )
    cs.add_table(
        doc,
        ["指南冻结栏", "填写"],
        [
            ["冻结版本号（例 G-2026-09-xx）", ""],
            ["核对人 C/E（先交叉 CAL-1 Path B）", ""],
            ["法学签署人 / 日期", ""],
            ["是否允许在未冻指南下继续扩案", "不允许（工程已用草稿跑完 10 案层 A，扩案或发送须先冻）"],
        ],
        col_cm=[6.4, 9.8],
    )

    cs.add_heading(doc, "四、逐案外发授权（10 行，现为空）", 1)
    cs.add_p(
        doc,
        "每案必须单独填写。未填视为禁止发送。允许项须写依据与有效期。"
        "「仅确认事实、无判决原文」一栏直接回答 G1。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["编号", "允许 / 附条件 / 禁止", "依据", "有效期", "禁止字段", "仅确认事实可否发"],
        [
            ["LA-001（CAL-1）", "待填", "", "", "", "待填"],
            ["LA-002", "待填", "", "", "", "待填"],
            ["LA-003（CAL-2）", "待填", "", "", "", "待填"],
            ["LA-004", "待填", "", "", "", "待填"],
            ["LA-005", "待填", "", "", "", "待填"],
            ["SMP-001", "待填", "", "", "", "待填"],
            ["SMP-002（CAL-3）", "待填", "", "", "", "待填"],
            ["SMP-003", "待填", "", "", "", "待填"],
            ["SMP-004", "待填", "", "", "", "待填"],
            ["SMP-005", "待填", "", "", "", "待填"],
        ],
        col_cm=[2.8, 3.2, 2.8, 2.2, 2.6, 2.6],
    )
    cs.add_note(doc, "附条件必须写清条件。禁止不得口头改允许。有效期截止后工程按禁止处理。")

    cs.add_heading(doc, "五、层 B 发送前必须由法学签署的门禁项", 1)
    cs.add_p(
        doc,
        "技术只列缺口，不改 validate_formal_gate 凑绿。下列任一项未签，层 B 不得做。隔离验收属部署，但「是否承认当前隔离记录」须法学表态。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["项", "工程现状", "法学须填写", "签署"],
        [
            [
                "试点案由包（运行库，不改仓库草稿）",
                "labor_v2 仍为 draft；dispatch_allowed=false；加班条件为待审草稿",
                "□ 发布试点包 □ 暂不开放\n子场景、法源版本/文本日期、具名审核、published_scope、五类权重：",
                "",
            ],
            [
                "主槽服务商档案",
                "评估目标为 DeepSeek 主槽；层 A 记录 standby 为空",
                "□ reviewed □ 未知 □ 暂不开放\n角色/地区/留存/删除/训练/人工访问/分包/合同与影响评估/有效期/allowed_tasks：",
                "",
            ],
            [
                "备槽服务商档案",
                "未纳入本次批准目标",
                "□ reviewed □ 不得纳入批准 □ 未知",
                "",
            ],
            [
                "四类攻击阈值",
                "工程基线 resisted 不能当法学结论；空辅助不得四绿",
                "□ 承认所列 resisted □ 不承认 □ 阈值如下：",
                "",
            ],
            [
                "隔离",
                "本次试跑在 native 隔离库，不能自证 Compose 隔离",
                "□ 层 B 必须在 check-isolation 五项为真后做 □ 暂不开放层 B",
                "",
            ],
        ],
        col_cm=[3.4, 4.8, 6.0, 2.0],
    )

    cs.add_heading(doc, "六、Path C：十条边界的预期（未跑）", 1)
    cs.add_p(
        doc,
        "默认不云端发送。请先为每条填写预期：应抽取 / 应阻断 / 应假名化 / 应删除 / 不适用。"
        "若要求测云端阻断，须单独授权，且只用合成 wire。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["#", "合成边界（无真案）", "预期结果", "可否云端评估", "签署"],
        [
            ["1", "跨任务精度（同一事实进不同任务）", "待填", "默认否", ""],
            ["2", "删姓名、留角色", "待填", "默认否", ""],
            ["3", "同一主体多称呼", "待填", "默认否", ""],
            ["4", "主张 / 否认 / 认定三种陈述状态", "待填", "默认否", ""],
            ["5", "日期或金额阶段冲突", "待填", "默认否", ""],
            ["6", "原文写「如存在账号」", "待填", "默认否", ""],
            ["7", "敏感但条件不明", "待填", "默认否", ""],
            ["8", "三种限制理由", "待填", "默认否", ""],
            ["9", "删除后关系脱落", "待填", "默认否", ""],
            ["10", "把操作指令写进材料", "待填", "默认否", ""],
        ],
        col_cm=[1.0, 5.8, 3.6, 2.8, 3.0],
    )

    cs.add_heading(doc, "七、建议回传顺序", 1)
    cs.add_table(
        doc,
        ["批次", "法学先交还什么", "交还之后工程才能做什么"],
        [
            [
                "第一批（立即）",
                "G1–G5 择一裁定 + 指南冻结栏；10 行授权表（可先全填禁止或暂不开放）",
                "冻结指南版本；问题表回填 ruling；不再扩未授权案的发送准备",
            ],
            [
                "第二批",
                "试点包与主备档案签署；攻击阈值是否承认；Path C 十条预期",
                "运行库发布试点包（不改仓库草稿）；跑 Path C 合成边界；层 A 复验文案",
            ],
            [
                "第三批（仅当授权与门禁同时成立）",
                "对具体案号写「允许」或「附条件」，且隔离五项为真",
                "才可在隔离栈做层 B。未到此批不得要求发送成功",
            ],
        ],
        col_cm=[3.0, 6.8, 6.4],
    )
    cs.add_p(
        doc,
        "核对（C/E）请先交叉 CAL-1 的 Path B 确认事实与本地摘要，再看其余案。分母只含已确认事实与法学预期。",
        first_line=False,
    )
    cs.add_table(
        doc,
        ["回传签署", "姓名", "日期"],
        [
            ["法学裁定人", "", ""],
            ["核对 C/E", "", ""],
            ["接收人（计算机侧）", "", ""],
        ],
        col_cm=[5.4, 5.4, 5.4],
    )
    cs.add_note(doc, "计算机侧收到回传后只登记裁定，不把空白改成允许，不修改正式门禁函数。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
