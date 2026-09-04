"""Generate only synthetic, reproducible demonstration attachments."""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docx import Document
from PIL import Image, ImageDraw
from lexaegis.demo import SCENARIOS

root = Path(__file__).resolve().parent.parent / "examples"
root.mkdir(exist_ok=True)
doc = Document()
doc.add_heading("纯合成劳动争议材料", 0)
doc.add_paragraph(SCENARIOS["success"]["note"])
for field, value in SCENARIOS["success"]["facts"]:
    doc.add_paragraph(value)
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "合成证据"
table.rows[0].cells[1].text = "说明"
row = table.add_row().cells
row[0].text, row[1].text = "工资表", "示例金额：12000元"
doc.save(root / "synthetic_labor.docx")
image = Image.new("RGB", (1000, 1400), "white")
draw = ImageDraw.Draw(image)
draw.text((90, 90), "SYNTHETIC SCANNED EVIDENCE - NOT A REAL CASE", fill="black", font_size=24)
draw.text((90, 150), "Demonstration amount: 12000 CNY", fill="black", font_size=22)
draw.text((90, 200), "Manual transcription or local OCR is required.", fill="black", font_size=22)
image.save(root / "synthetic_scan.pdf", "PDF", resolution=100)
print("Created synthetic DOCX and scanned PDF in examples/")
