"""Import real-case sensitivity classification workbooks into machine-readable JSON.

Reads each case's ``case.json`` under ``.data/real_cases/cases/`` and resolves the
``role: sensitivity`` material (either ``raw`` xlsx or shared ``workbook``+``sheet``).
Rows are written verbatim (original Chinese headers as keys) to
``working/sensitivity_classification.json`` with provenance, and the material's
``working`` field is linked back in ``case.json``.

All output stays under ``.data/`` (gitignored). Sources are never modified.
"""
from __future__ import annotations

import argparse
import datetime
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

CASES_DIR = Path(".data/real_cases/cases")
OUT_NAME = "sensitivity_classification.json"
MANIFEST_PATH = Path(".data/real_cases/annotation_import_manifest.json")
ANNOTATION_METHOD = "dual_annotator_merged_final"
ANNOTATORS = ["林语", "叶云"]
ADJUDICATOR = "苏苒"


def sha256_of(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_sheet(path: Path, sheet: str | None) -> tuple[list[str], list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ValueError(f"{path} sheet {sheet!r}: no rows")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    body = []
    for idx, raw in enumerate(rows[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in raw):
            continue
        body.append({"row": idx, **{h: (None if v is None else str(v).strip()) for h, v in zip(header, raw) if h}})
    return [h for h in header if h], body


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write", action="store_true", help="write JSON output (default: dry-run report only)")
    args = parser.parse_args()

    today = datetime.date.today().isoformat()
    manifest = {"schema_version": 1, "kind": "sensitivity_classification_import",
                "annotation_method": ANNOTATION_METHOD, "annotators": ANNOTATORS,
                "adjudicator": ADJUDICATOR, "imported_on": today, "cases": []}
    errors = []

    for case_dir in sorted(CASES_DIR.iterdir()):
        case_json_path = case_dir / "case.json"
        if not case_json_path.exists():
            continue
        case = json.loads(case_json_path.read_text(encoding="utf-8"))
        case_id = case.get("id", case_dir.name)
        material = next((m for m in case.get("materials", []) if m.get("role") == "sensitivity"), None)
        if not material:
            errors.append(f"{case_id}: no sensitivity material registered")
            continue

        if material.get("raw"):
            source = case_dir / material["raw"]
            sheet = None
        elif material.get("workbook"):
            source = CASES_DIR.parent / material["workbook"]
            sheet = material["sheet"]
        else:
            errors.append(f"{case_id}: sensitivity material has no raw/workbook path")
            continue
        if not source.exists():
            errors.append(f"{case_id}: source missing {source}")
            continue

        digest = sha256_of(source)
        registered = material.get("sha256")
        hash_ok = registered in (None, digest)

        try:
            columns, rows = read_sheet(source, sheet)
        except Exception as exc:
            errors.append(f"{case_id}: parse failed: {exc}")
            continue

        record = {
            "schema_version": 1,
            "case_id": case_id,
            "kind": "sensitivity_classification",
            "annotation_method": ANNOTATION_METHOD,
            "annotators": ANNOTATORS,
            "adjudicator": ADJUDICATOR,
            "provenance": {
                "source_path": str(source),
                "source_sha256": digest,
                "registered_sha256": registered,
                "sha256_match": hash_ok,
                "workbook_sheet": sheet,
                "original_name": material.get("original_name"),
                "imported_on": today,
                "importer": "scripts/import_case_sensitivity.py",
            },
            "columns": columns,
            "row_count": len(rows),
            "rows": rows,
        }
        manifest["cases"].append({
            "case_id": case_id, "row_count": len(rows), "source": str(source),
            "sheet": sheet, "sha256_match": hash_ok,
            "output": f"cases/{case_id}/working/{OUT_NAME}",
        })
        status = "ok" if hash_ok else "HASH-MISMATCH"
        print(f"{case_id}: {len(rows):3d} rows  [{status}]  {source.name}" + (f"#{sheet}" if sheet else ""))

        if args.write:
            out_path = case_dir / "working" / OUT_NAME
            out_path.parent.mkdir(parents=True, exist_ok=True)
            out_path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
            material["working"] = f"working/{OUT_NAME}"
            case_json_path.write_text(json.dumps(case, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    for e in errors:
        print(f"ERROR: {e}", file=sys.stderr)
    if args.write and manifest["cases"]:
        MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"\nmanifest -> {MANIFEST_PATH}")
    if not args.write:
        print("\ndry-run: pass --write to emit JSON")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
